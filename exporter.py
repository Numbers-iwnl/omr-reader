\
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import List

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from omr_reader import Answer


def export_answers_to_xlsx(answers: List[Answer], image_path: str, out_path: str) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Respostas"

    ws["A1"] = "Arquivo de origem"
    ws["B1"] = Path(image_path).name
    ws["A2"] = "Gerado em"
    ws["B2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    ws.append([])  # linha em branco

    ws.append(["Questão", "Resposta", "Status", "Score (média de cinza)"])

    for a in answers:
        ws.append([a.question, a.answer, a.status, round(float(a.score), 2)])

    # Ajuste de largura
    for col in range(1, 5):
        ws.column_dimensions[get_column_letter(col)].width = [10, 10, 14, 18][col - 1]

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return str(out)
